"""Excel export.

The workbook is the deliverable, so it has to carry the uncertainty with the
data rather than beside it. Three rules shape the layout:

  * Every row states how it was arrived at -- how many frames saw the book,
    which frames, and what the automated checks said about it.
  * The rows that need a person are a separate sheet, not a colour buried in
    a thousand-row table. That sheet IS the work queue.
  * Counts are live formulas over the Catalogue sheet, so correcting a row by
    hand updates the summary instead of silently contradicting it.
"""
import datetime as _dt

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

FONT = "Arial"
HDR_FILL = PatternFill("solid", fgColor="1F3864")
BAND_FILL = PatternFill("solid", fgColor="F2F5FA")
CONF_FILL = {"high": PatternFill("solid", fgColor="E2EFE9"),
             "medium": PatternFill("solid", fgColor="FBF3DF"),
             "low": PatternFill("solid", fgColor="FBE5E4"),
             "none": PatternFill("solid", fgColor="E8E8E8")}
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def confidence(book):
    """Two independent axes decide this: could the spine be read, and was the
    read confirmed by a second frame. Neither alone is enough -- a crisp read
    seen once can still be a misread, and a corroborated blur is still a blur."""
    leg = book.legibility
    corroborated = book.n_reads >= 2 and "read_disagreement" not in book.flags
    if "read_disagreement" in book.flags:
        return "low"
    if leg == "illegible":
        return "none"
    if leg == "clear":
        return "high" if corroborated else "medium"
    return "medium" if corroborated else "low"


def _write_row(ws, r, values):
    """Explicit row writes, never ws.append().

    append() targets ws.max_row + 1, and simply *reading* a cell materialises
    its row -- so setting freeze_panes (which references row 2) silently
    pushes the first appended row to 3 and leaves a blank row in the data.
    That blank row then breaks every COUNTA over the sheet."""
    for c, v in enumerate(values, start=1):
        ws.cell(row=r, column=c, value=v)


def _style_header(ws, ncols, row=1):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = Font(name=FONT, bold=True, color="FFFFFF", size=10)
        cell.fill = HDR_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = ws.cell(row=row + 1, column=1)
    ws.auto_filter.ref = f"A{row}:{get_column_letter(ncols)}{row}"


def _widths(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


CAT_COLS = ["#", "Shelf", "Row", "Position", "Title", "Author / editors", "Volume",
            "Detail", "Publisher", "Confidence", "Legibility", "Frames seen",
            "Source images", "Flags", "Notes"]
# Column letters used by the Summary formulas. Kept next to CAT_COLS so a
# column insertion cannot silently point a COUNTIF at the wrong data.
COL = {name: chr(ord("A") + i) for i, name in enumerate(CAT_COLS)}


def build(path, layers, frames_meta, anomalies, duplicates, suspect_joins, joins,
          project="UATX library", assignments=None, location_warnings=None,
          row_decisions=None):
    wb = Workbook()

    # ---------- Catalogue ----------
    ws = wb.active
    ws.title = "Catalogue"
    _write_row(ws, 1, CAT_COLS)
    _style_header(ws, len(CAT_COLS))
    row = 2
    index = []
    assignments = assignments or [{"shelf": i + 1, "row": "-"} for i in range(len(layers))]
    for li, layer in enumerate(layers, start=1):
        loc = assignments[li - 1]
        for pos, b in enumerate(layer, start=1):
            conf = confidence(b)
            note = "; ".join(n for n in [getattr(b, "note", None)] if n)
            _write_row(ws, row, [row - 1, loc["shelf"], loc["row"], pos, b.title or "", b.author or "",
                                 b.volume or "", getattr(b, "detail", "") or "",
                                 getattr(b, "publisher", "") or "", conf, b.legibility,
                                 b.n_reads, ", ".join(sorted(set(b.sources))),
                                 ", ".join(b.flags), note])
            for c in range(1, len(CAT_COLS) + 1):
                cell = ws.cell(row=row, column=c)
                cell.font = Font(name=FONT, size=10)
                cell.alignment = Alignment(vertical="top", wrap_text=(c in (5, 6, 8, 13, 14, 15)))
                cell.border = BORDER
            ci = CAT_COLS.index("Confidence") + 1
            ws.cell(row=row, column=ci).fill = CONF_FILL[conf]
            ws.cell(row=row, column=ci).font = Font(name=FONT, size=10, bold=True)
            index.append((row, li, pos, b, conf, loc))
            row += 1
    _widths(ws, [5, 7, 9, 8, 40, 30, 12, 30, 11, 12, 11, 7, 22, 30, 34])
    last = row - 1

    # ---------- Review queue ----------
    rq = wb.create_sheet("Review queue")
    _write_row(rq, 1, ["Priority", "#", "Shelf", "Row", "Position", "What to check",
                       "Title", "Volume", "Confidence", "Source images"])
    _style_header(rq, 10)
    r = 2
    queue = []
    for (xrow, li, pos, b, conf, loc) in index:
        reasons = []
        if conf == "none":
            reasons.append("spine unreadable - pull the book")
        if "read_disagreement" in b.flags:
            reasons.append(f"frames disagree: {' / '.join(b.variants)}")
        if conf == "low" and "read_disagreement" not in b.flags:
            reasons.append("partial read, seen only once")
        # A note is only a review trigger when the row is not already high
        # confidence. 461 of 925 rows carry a transcriber note; queueing all of
        # them buries the 238 that genuinely need a person.
        if getattr(b, "note", None) and conf != "high":
            reasons.append(b.note)
        if reasons:
            pri = 1 if conf == "none" or "read_disagreement" in b.flags else 2
            queue.append((pri, xrow, li, pos, "; ".join(reasons), b, conf, loc))
    for pri, xrow, li, pos, why, b, conf, loc in sorted(queue, key=lambda q: (q[0], q[1])):
        _write_row(rq, r, [pri, xrow - 1, loc["shelf"], loc["row"], pos, why, b.title or "",
                           b.volume or "", conf, ", ".join(sorted(set(b.sources)))])
        for c in range(1, 11):
            rq.cell(row=r, column=c).font = Font(name=FONT, size=10)
            rq.cell(row=r, column=c).alignment = Alignment(vertical="top", wrap_text=(c in (6, 7)))
            rq.cell(row=r, column=c).border = BORDER
        r += 1
    review_last = r - 1
    _widths(rq, [9, 5, 7, 9, 8, 52, 38, 12, 12, 22])

    # ---------- Anomalies ----------
    an = wb.create_sheet("Anomalies")
    _write_row(an, 1, ["Type", "Shelf layer", "Position", "Detail", "Meaning"])
    _style_header(an, 5)
    r = 2
    MEAN = {"out_of_order": "Shelved out of sequence, or the volume number was misread.",
            "gap": "The library appears not to hold these volumes. Confirm before assuming loss.",
            "duplicate": "A second physical copy. Legitimate - recorded, not removed.",
            "suspect_join": "Two layers may actually be one; check for repeated books."}
    def _loc(li):
        a = assignments[li - 1]
        return f"Shelf {a['shelf']} {a['row']}"
    for li, items in anomalies:
        for a in items:
            _write_row(an, r, [a["kind"], _loc(li),
                               a["at"] + 1 if isinstance(a.get("at"), int) else "",
                               a["detail"], MEAN.get(a["kind"], "")])
            r += 1
    for d in duplicates:
        _write_row(an, r, ["duplicate", _loc(int(d["a"].split("#")[0][1:]) + 1), "",
                           f"{d['title_a'][:60]} ({d.get('volume') or 'no volume'}) appears twice",
                           MEAN["duplicate"]])
        r += 1
    for s in suspect_joins:
        _write_row(an, r, ["suspect_join", f"{_loc(s['left_layer']+1)} / {_loc(s['right_layer']+1)}", "",
                           f"{s['shared']} book(s) shared at {s['similarity']}% similarity",
                           MEAN["suspect_join"]])
        r += 1
    for d in (row_decisions or []):
        if not d["verified"]:
            _write_row(an, r, ["row_seam_assumed", "", "",
                               f"{d['between']}: {d['evidence']}",
                               "These frames were joined into one row on the absence of "
                               "a shelf-end gap. Book order across this seam follows photo "
                               "order and is not verified by shared books."])
            r += 1
    for w in (location_warnings or []):
        _write_row(an, r, [f"location/{w['severity']}", "", "", w["detail"],
                           "Affects which shelf and row books are filed under, "
                           "not whether they were read correctly."])
        r += 1
    for rr in range(2, r):
        for c in range(1, 6):
            an.cell(row=rr, column=c).font = Font(name=FONT, size=10)
            an.cell(row=rr, column=c).alignment = Alignment(vertical="top", wrap_text=(c in (4, 5)))
            an.cell(row=rr, column=c).border = BORDER
    _widths(an, [15, 14, 10, 56, 56])

    # ---------- Frames ----------
    fr = wb.create_sheet("Frames")
    _write_row(fr, 1, ["Frame", "Captured", "Pixels", "Sharpness", "Quality",
                       "Band y0", "Band y1", "Note"])
    _style_header(fr, 8)
    for i, f in enumerate(frames_meta, start=2):
        _write_row(fr, i, [f.get("frame") or f.get("name", ""), f.get("captured_at", "") or "", f.get("size", ""),
                           f.get("sharpness", ""), f.get("quality", ""),
                           f.get("y0f", ""), f.get("y1f", ""), f.get("quality_note", "")])
        for c in range(1, 9):
            fr.cell(row=i, column=c).font = Font(name=FONT, size=10)
            fr.cell(row=i, column=c).border = BORDER
    _widths(fr, [14, 20, 14, 11, 10, 10, 10, 40])

    # ---------- Summary ----------
    sm = wb.create_sheet("Summary", 0)
    sm["A1"] = f"{project} - spine catalogue"
    sm["A1"].font = Font(name=FONT, size=15, bold=True, color="1F3864")
    sm["A2"] = f"Generated {_dt.datetime.now():%d %B %Y %H:%M} from {len(frames_meta)} photographs"
    sm["A2"].font = Font(name=FONT, size=10, italic=True, color="595959")

    C, K = COL["Confidence"], COL["Frames seen"]
    rows = [("Volumes catalogued", f'=COUNTA(Catalogue!A2:A{last})'),
            ("  high confidence", f'=COUNTIF(Catalogue!{C}2:{C}{last},"high")'),
            ("  medium confidence", f'=COUNTIF(Catalogue!{C}2:{C}{last},"medium")'),
            ("  low confidence", f'=COUNTIF(Catalogue!{C}2:{C}{last},"low")'),
            ("  unreadable", f'=COUNTIF(Catalogue!{C}2:{C}{last},"none")'),
            ("Seen in 2+ frames (corroborated)", f'=COUNTIF(Catalogue!{K}2:{K}{last},">=2")'),
            ("Seen once only", f'=COUNTIF(Catalogue!{K}2:{K}{last},1)'),
            ("Rows in review queue", f'=COUNTA(\'Review queue\'!B2:B{max(2, review_last)})'),
            ("Shelves", len({a["shelf"] for a in assignments})),
            ("Rows (shelf layers)", len(layers)),
            ("Photographs ingested", len(frames_meta))]
    sm["A4"] = "Counts"; sm["A4"].font = Font(name=FONT, bold=True, size=11)
    rr = 5
    for label, val in rows:
        sm.cell(row=rr, column=1, value=label).font = Font(name=FONT, size=10,
                                                           bold=not label.startswith("  "))
        c = sm.cell(row=rr, column=2, value=val)
        c.font = Font(name=FONT, size=10); c.alignment = Alignment(horizontal="right")
        rr += 1

    rr += 1
    sm.cell(row=rr, column=1, value="How to read this workbook").font = Font(name=FONT, bold=True, size=11)
    rr += 1
    for line in [
        "Catalogue - every spine seen, one row each. Nothing is dropped: an unreadable spine still gets a row with its Shelf, Row and Position so it can be filled in later.",
        "Location: Shelf is the bookcase (1, 2, 3 ...); Row is top / middle / bottom within it; Position counts left to right along that row.",
        "Review queue - the rows a person needs to resolve, most urgent first. This sheet is the work list.",
        "Anomalies - volume-sequence problems, duplicate copies and uncertain layer boundaries found by automated checks.",
        "Frames - one row per photograph, with the sharpness and shelf band used. This is the audit trail back to the pixels.",
        "",
        "Confidence is derived, not guessed: high = clearly legible AND read in two or more overlapping frames;",
        "medium = clear but seen once, or partial but corroborated; low = partial and seen once, or frames disagreed;",
        "none = the spine could not be read at all.",
        "",
        "Titles and authors are transcribed from the spine only. No publication year or ISBN is recorded, because",
        "neither can be read from a spine - supplying them would mean inventing them.",
    ]:
        sm.cell(row=rr, column=1, value=line).font = Font(name=FONT, size=9,
                                                          italic=line.startswith("Confidence") or not line)
        rr += 1
    _widths(sm, [88, 14])

    wb.save(path)
    return {"rows": last - 1, "review": review_last - 1, "layers": len(layers)}
